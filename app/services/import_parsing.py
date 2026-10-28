# app/services/import_parsing.py
"""
Reading import files and understanding the values in them (C-05, C-06). No database access.

read_table(): a CSV or XLSX as header names + rows of text cells.
  - CSV: UTF-8 or Windows-1252 (app/utils/upload_guards.py), delimiter sniffed (, ; tab |).
  - XLSX: zip-bomb guard first, then openpyxl read-only; the person may pick the sheet (the first sheet
    with data otherwise). Cells become text on purpose: phone numbers and postcodes must not turn into
    numbers; whole floats lose Excel's ".0"; dates become ISO dates.
  - Header names are trimmed, blanks named "Column 3", duplicates numbered "Email (2)".
  - Size, row and column limits (ImportLimits) are checked while reading.

Value parsers (strict, with a reason for the row when a value cannot be understood):
  dates (YYYY-MM-DD, DMY / MDY with a per-import date order, "3 Mar 2026", Excel serial numbers),
  numbers (thousands separators, currency signs, (1,234) negatives), Yes / No / true / false,
  choices matched by value or label ignoring case, "Jane Doe <jane@x.com>" reduced to the address,
  "Doe, Jane" split into first / last name, Excel's ".0" stripped from phone numbers.
"""

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import pandas as pd
from openpyxl import load_workbook

from app.domain.imports import DateOrder, ImportLimits
from app.utils.upload_guards import FileRejected, check_column_count, check_size, check_xlsx_archive, decode_csv

Row = Dict[str, Optional[str]]

TRUE_WORDS = frozenset({"true", "t", "yes", "y", "1", "x", "checked", "on"})
FALSE_WORDS = frozenset({"false", "f", "no", "n", "0", "unchecked", "off"})
_EXCEL_EPOCH = date(1899, 12, 30)
_DATE_FORMATS = ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%d %b %Y", "%d %B %Y", "%b %d, %Y", "%B %d, %Y",
                 "%d-%b-%Y", "%d-%B-%Y", "%b %d %Y", "%B %d %Y", "%d %b, %Y")
_NUMERIC_DATE = re.compile(r"(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{2}|\d{4})")
_TIME_SUFFIX = re.compile(r"[T ]\d{1,2}:\d{2}(:\d{2}(\.\d+)?)?\s*([AaPp][Mm])?Z?$")
_ANGLE_EMAIL = re.compile(r"<\s*([^<>@\s]+@[^<>\s]+?)\s*>")
_MAILTO = re.compile(r"^mailto:", re.I)
_EXCEL_FLOAT = re.compile(r"^(\+?\d+)\.0+$")


class ValueProblem(ValueError):
    """A cell value that cannot be read; the message names the field and the value."""


# ── File reading ──────────────────────────────────────────────────

@dataclass
class Table:
    headers: List[str]
    rows: List[Row]
    sheet_names: List[str] = field(default_factory=list)
    sheet: Optional[str] = None


def clean_headers(raw: Sequence[Any]) -> List[str]:
    """Trim, fill blanks ("Column 3") and number duplicates ("Email (2)")."""
    seen: Dict[str, int] = {}
    out = []
    for i, h in enumerate(raw, 1):
        name = " ".join(str("" if h is None else h).replace("﻿", "").split()) or f"Column {i}"
        n = seen.get(name.lower(), 0) + 1
        seen[name.lower()] = n
        out.append(name if n == 1 else f"{name} ({n})")
    return out


def cell_text(value: Any) -> Optional[str]:
    """One spreadsheet cell as text (None when blank)."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0, 0) else value.replace(microsecond=0).isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        text = str(int(value)) if value.is_integer() and abs(value) < 1e15 else format(value, ".15g")
    else:
        text = str(value).strip()
    return text or None


def _delimiter(text: str) -> str:
    """Comma, semicolon, tab or pipe; a one-column file has none, so default to comma."""
    try:
        return csv.Sniffer().sniff(text[:20000], delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def _grid_csv(raw: bytes) -> List[List[Any]]:
    text = decode_csv(raw)
    try:
        return [cells for cells in csv.reader(io.StringIO(text), delimiter=_delimiter(text))]
    except csv.Error as exc:
        raise FileRejected(f"The CSV could not be read ({exc}).") from exc


def _sheet_rows(sheet: Any, max_rows: int) -> List[List[Any]]:
    grid: List[List[Any]] = []
    for values in sheet.iter_rows(values_only=True):
        grid.append(list(values))
        if len(grid) > max_rows + 1 and any(v not in (None, "") for v in values):
            break  # enough to know it is over the limit
    return grid


def _grid_xlsx(raw: bytes, sheet: Optional[str], max_rows: int) -> Tuple[List[List[Any]], List[str], Optional[str]]:
    check_xlsx_archive(raw)
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise FileRejected("The workbook could not be opened. Save it again as .xlsx or .csv and retry.") from exc
    try:
        names = list(workbook.sheetnames)
        if sheet and sheet not in names:
            raise FileRejected(f"Sheet '{sheet}' was not found in the workbook.")
        for name in [sheet] if sheet else names:
            grid = _sheet_rows(workbook[name], max_rows)
            if any(any(v not in (None, "") for v in row) for row in grid):
                return grid, names, name
        return [], names, sheet or (names[0] if names else None)
    finally:
        workbook.close()


def _grid_xls(raw: bytes) -> List[List[Any]]:
    frame = pd.read_excel(io.BytesIO(raw), dtype=str, header=None)
    return frame.astype(object).where(pd.notna(frame), None).values.tolist()


def _to_table(grid: List[List[Any]], limits: ImportLimits) -> Tuple[List[str], List[Row]]:
    rows_iter = iter(row for row in grid if any(cell_text(v) for v in row))
    header = next(rows_iter, None)
    if header is None:
        raise FileRejected("The file is empty.")
    width = len(header)
    while width and cell_text(header[width - 1]) is None:  # ignore empty trailing columns
        width -= 1
    if width == 0:
        raise FileRejected("The first row must contain the column names.")
    check_column_count(range(width), limits.max_columns)
    headers = clean_headers(header[:width])
    rows: List[Row] = []
    for cells in rows_iter:
        padded = list(cells[:width]) + [None] * (width - len(cells))
        rows.append({h: cell_text(v) for h, v in zip(headers, padded)})
        if len(rows) > limits.max_rows:
            raise FileRejected(f"Files can have at most {limits.max_rows:,} rows; split this one into smaller files.")
    return headers, rows


def read_table(raw: bytes, file_name: str, limits: ImportLimits, sheet: Optional[str] = None) -> Table:
    """The file as a table of text cells. Raises FileRejected with a message for the person."""
    if not raw:
        raise FileRejected("The file is empty.")
    check_size(raw, limits.max_bytes)
    name = (file_name or "").lower()
    names: List[str] = []
    used: Optional[str] = None
    if name.endswith((".xlsx", ".xlsm")):
        grid, names, used = _grid_xlsx(raw, sheet, limits.max_rows)
    elif name.endswith(".xls"):
        grid = _grid_xls(raw)
    else:
        grid = _grid_csv(raw)
    headers, rows = _to_table(grid, limits)
    return Table(headers, rows, names, used)


def table_frame(table: Table) -> pd.DataFrame:
    """A table as a pandas DataFrame of text cells (None for blanks)."""
    return pd.DataFrame(table.rows, columns=table.headers)


# ── Values ─────────────────────────────────────────────────────────

def _year(y: int) -> int:
    return y + (2000 if y < 70 else 1900) if y < 100 else y


def _numeric_date(a: int, b: int, y: int, order: DateOrder) -> date:
    if a > 12 >= b:
        day, month = a, b
    elif b > 12 >= a:
        day, month = b, a
    elif a <= 12 and b <= 12:
        day, month = (a, b) if order is DateOrder.DMY else (b, a)
    else:
        raise ValueError("not a date")
    return date(_year(y), month, day)


def parse_date(text: str, order: DateOrder = DateOrder.MDY) -> date:
    """A date from the ways people write them; raises ValueError."""
    t = " ".join(str(text).strip().split())
    if re.fullmatch(r"\d{5}(\.\d+)?", t) and 20000 <= float(t) <= 80000:  # Excel serial (1954-2119)
        return _EXCEL_EPOCH + timedelta(days=int(float(t)))
    t = _TIME_SUFFIX.sub("", t).strip()
    m = _NUMERIC_DATE.fullmatch(t)
    if m:
        return _numeric_date(int(m.group(1)), int(m.group(2)), int(m.group(3)), order)
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(t, fmt).date()
        except ValueError:
            continue
    raise ValueError("not a date")


def parse_number(text: str) -> float:
    """1,234.50  $1,200  (1,234)  -5  12%  -> float; raises ValueError."""
    t = str(text).strip()
    negative = t.startswith("(") and t.endswith(")")
    cleaned = re.sub(r"[^0-9.\-]", "", t.replace(",", ""))
    if cleaned in ("", "-", ".", "-.") or cleaned.count(".") > 1 or "-" in cleaned[1:]:
        raise ValueError("not a number")
    number = float(cleaned)
    return -abs(number) if negative else number


def parse_bool(text: str) -> bool:
    word = str(text).strip().lower()
    if word in TRUE_WORDS:
        return True
    if word in FALSE_WORDS:
        return False
    raise ValueError("not yes / no")


def _slug(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value).lower())


def _option_pairs(options: Iterable[Any]) -> List[Tuple[str, str]]:
    pairs = []
    for option in options or []:
        if isinstance(option, dict):
            value = str(option.get("value", ""))
            pairs.append((value, str(option.get("label", value))))
        else:
            pairs.append((str(option), str(option)))
    return pairs


def match_choice(options: Iterable[Any], text: str) -> Optional[str]:
    """The stored option value for text, matched on value or label ignoring case and punctuation.
    Yes / No words also match a Yes / No (true / false) option list."""
    pairs = _option_pairs(options)
    wanted = _slug(text)
    for value, label in pairs:
        if wanted and wanted in (_slug(value), _slug(label)):
            return value
    try:
        flag = parse_bool(text)
    except ValueError:
        return None
    for value, label in pairs:
        try:
            if parse_bool(value) is flag or parse_bool(label) is flag:
                return value
        except ValueError:
            continue
    return None


def email_address(text: Optional[str]) -> Optional[str]:
    """"Jane Doe <jane@x.com>" / "mailto:jane@x.com" -> "jane@x.com"; other text trimmed."""
    if text is None:
        return None
    t = str(text).strip()
    m = _ANGLE_EMAIL.search(t)
    if m:
        t = m.group(1)
    return _MAILTO.sub("", t).strip() or None


def split_name(text: str) -> Tuple[Optional[str], Optional[str]]:
    """"Jane van der Berg" -> ("Jane", "van der Berg"); "Doe, Jane" -> ("Jane", "Doe")."""
    t = " ".join((text or "").split())
    if not t:
        return None, None
    if "," in t:
        last, _, first = t.partition(",")
        return first.strip() or None, last.strip() or None
    first, _, last = t.partition(" ")
    return first or None, last.strip() or None


def clean_phone(text: Optional[str]) -> Optional[str]:
    """Excel turns 4155550100 into "4155550100.0"; the ".0" is not part of the number."""
    if text is None:
        return None
    t = str(text).strip()
    m = _EXCEL_FLOAT.match(t)
    return (m.group(1) if m else t) or None


def custom_value(field_type: str, label: str, options: Iterable[Any], text: str, order: DateOrder) -> Any:
    """A cell for a custom property in the form app/services/contact_field_service accepts."""
    try:
        if field_type == "NUMBER":
            number = parse_number(text)
            return int(number) if number.is_integer() else number
        if field_type == "DATE":
            return parse_date(text, order).isoformat()
        if field_type == "PHONE":
            return clean_phone(text)
        if field_type in ("SELECT", "RADIO"):
            return match_choice(options, text) or text
        if field_type == "MULTI_CHECKBOX":
            parts = [p.strip() for p in re.split(r"[;|\n]+", text) if p.strip()]
            return [match_choice(options, p) or p for p in parts]
    except ValueError:
        what = {"NUMBER": "a number", "DATE": "a date"}.get(field_type, "valid")
        raise ValueProblem(f"{label}: '{text}' is not {what}") from None
    return text


def data_error_reason(exc: BaseException, labels: Dict[str, str]) -> str:
    """'Job title: value is too long ...' from MySQL's "Data too long for column 'designation'"."""
    match = re.search(r"column '([^']+)'", str(getattr(exc, "orig", exc)))
    if match:
        column = match.group(1)
        return f"{labels.get(column, column)}: value is too long or not valid for this field"
    return "A value is too long or not valid for its field"
