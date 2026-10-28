"""
Word layout of the analytics report (GET /reports/export/global-analytics/docx), split out of
app/services/reports_service.py (chunk A6). Pure document building: the service passes the rows.
"""

import io
from datetime import date, datetime
from typing import Any, List, Optional, Sequence

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BLUE = RGBColor(0x2D, 0x6B, 0xBF)
GREY = RGBColor(0x6B, 0x72, 0x80)
LIGHT = RGBColor(0xA0, 0xAA, 0xB5)
INK = RGBColor(0x11, 0x18, 0x27)

CAMPAIGN_HEADERS = ["Campaign Name", "Status", "Sent", "Opened", "Replied", "Bounced", "Open Rate %",
                    "Reply Rate %", "Bounce Rate %", "Created"]
MAILBOX_HEADERS = ["Mailbox", "Leads Contacted", "Emails Sent", "Opened", "Open Rate %", "Replied", "Reply Rate %",
                   "Bounced", "Bounce Rate %"]
PROVIDER_HEADERS = ["Provider", "Sent", "Opened", "Replied", "Bounced", "Open Rate %", "Reply Rate %",
                    "Bounce Rate %"]


def _set_cell_bg(cell: Any, hex_color: str) -> None:
    """Fill a table cell background with a hex color (e.g. '2d6bbf')."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _style(run: Any, size: float, color: RGBColor, bold: bool = False) -> None:
    run.bold = bold or None
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    run.font.color.rgb = color


def _text_paragraph(doc: DocxDocument, text: str, size: float, color: RGBColor, space_after: Optional[float],
                    bold: bool = False) -> Any:
    para = doc.add_paragraph()
    if space_after is not None:
        para.paragraph_format.space_after = Pt(space_after)
    _style(para.add_run(text), size, color, bold)
    return para


def _section_heading(doc: DocxDocument, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    _style(p.add_run(text.upper()), 9, BLUE, bold=True)
    p_pr = p._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:color"), "2d6bbf")
    border.append(bottom)
    p_pr.append(border)


def _kpi_table(doc: DocxDocument, rows: Sequence[tuple]) -> None:
    """Two-column key-value table for the KPI summary."""
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, (label, value) in enumerate(rows):
        lc, vc = table.cell(i, 0), table.cell(i, 1)
        lc.text, vc.text = label, str(value)
        _style(lc.paragraphs[0].runs[0], 10, GREY, bold=True)
        _style(vc.paragraphs[0].runs[0], 10, INK, bold=True)
        bg = "F9FAFB" if i % 2 == 0 else "FFFFFF"
        _set_cell_bg(lc, bg)
        _set_cell_bg(vc, bg)
    for row in table.rows:
        row.cells[0].width = Inches(2.8)
        row.cells[1].width = Inches(1.5)


def _data_table(doc: DocxDocument, headers: List[str], data_rows: Sequence[Sequence[Any]]) -> None:
    """Full data table with a blue header row."""
    table = doc.add_table(rows=1 + len(data_rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        _style(cell.paragraphs[0].runs[0], 8, RGBColor(0xFF, 0xFF, 0xFF), bold=True)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_cell_bg(cell, "2D6BBF")
    for r_idx, row_data in enumerate(data_rows):
        cells = table.rows[r_idx + 1].cells
        bg = "F9FAFB" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cells[c_idx].text = str(val) if val is not None else ""
            _style(cells[c_idx].paragraphs[0].runs[0], 8, INK)
            _set_cell_bg(cells[c_idx], bg)


def analytics_report(*, start_date: date, end_date: date, user_label: str, kpi_rows: Sequence[tuple],
                     timeseries: Sequence[Sequence[Any]], campaigns: Sequence[Sequence[Any]],
                     mailboxes: Sequence[Sequence[Any]], providers: Sequence[Sequence[Any]]) -> io.BytesIO:
    """Cover header, then 1. KPI summary, 2. daily series, 3. per campaign, 4. mailbox health,
    5. provider performance."""
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(1)

    _text_paragraph(doc, "OUTREACH360 — ANALYTICS REPORT", 18, BLUE, 4, bold=True)
    _text_paragraph(doc, f"Period: {start_date.strftime('%B %d, %Y')} – {end_date.strftime('%B %d, %Y')}", 10, GREY, 2)
    _text_paragraph(doc, f"Filtered By: {user_label}", 10, GREY, 2)
    # Local server time, as the report has always printed it
    _text_paragraph(doc, f"Generated: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}", 9, LIGHT, 16)

    _section_heading(doc, "1. Overall KPI Summary")
    _kpi_table(doc, kpi_rows)
    _section_heading(doc, "2. Daily Sending Activity")
    _data_table(doc, ["Date", "Sent", "Opened", "Replied", "Bounced"], timeseries)
    doc.add_page_break()
    _section_heading(doc, "3. Per-Campaign Breakdown")
    _data_table(doc, CAMPAIGN_HEADERS, campaigns)
    _section_heading(doc, "4. Mailbox Health")
    _data_table(doc, MAILBOX_HEADERS, mailboxes)
    _section_heading(doc, "5. Provider Performance")
    _data_table(doc, PROVIDER_HEADERS, providers)

    doc.add_paragraph()
    footer = _text_paragraph(doc, "This report was automatically generated by Outreach360.ai", 8, LIGHT, None)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
