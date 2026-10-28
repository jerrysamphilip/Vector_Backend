"""
A7 functional tests (MySQL): background CRM imports (C-03, F-16), options (C-04), value parsing (C-05) and
file reading (C-06).

  * upload -> dry run (writes nothing) -> start -> worker pass in chunks -> COMPLETED with parsed values
  * a worker that died mid-way: the stale RUNNING job is resumed at its checkpoint, not re-imported
  * cancel between chunks; cancel before it runs; resume of a failed job
  * modes CREATE / UPDATE_ONLY on the synchronous endpoint; a file over the sync limit is queued (202)
  * saved mappings; tenant isolation; manage_prospects required
"""
import io
import json
import uuid

import pytest
from openpyxl import Workbook

from app.adapters.storage import LocalFileStore
from app.services import import_service
from tests.conftest import RUN, SessionLocal, auth, login, register_tenant, sql


@pytest.fixture(scope="module")
def ws(client):
    return register_tenant(client, "aimp")


@pytest.fixture(scope="module")
def other(client):
    return register_tenant(client, "aimq")


@pytest.fixture(autouse=True)
def _store(tmp_path, monkeypatch):
    store = LocalFileStore(tmp_path / "imports")
    monkeypatch.setattr(import_service, "import_store", lambda: store)
    return store


def _uid():
    return uuid.uuid4().hex[:6]


def _email(label):
    return f"{label}.{_uid()}.{RUN}@aimp{RUN}.example.com"


def _upload(client, h, content: bytes, name="contacts.csv", **data):
    r = client.post("/imports/uploads", headers=h, files={"file": (name, io.BytesIO(content), "text/csv")}, data=data)
    assert r.status_code == 201, r.text
    return r.json()


def _work(**kw):
    return import_service.run_background(SessionLocal, "test-worker", **kw)


def _job(client, h, import_id):
    r = client.get(f"/imports/{import_id}", headers=h)
    assert r.status_code == 200, r.text
    return r.json()


def test_background_import_end_to_end(client, ws):
    h = ws["headers"]
    r = client.post("/contacts/fields", headers=h, json={"label": f"Seats {_uid()}", "field_type": "NUMBER"})
    assert r.status_code == 201, r.text
    seats = r.json()["field_key"]
    r = client.post("/contacts/fields", headers=h, json={"label": f"Renewal {_uid()}", "field_type": "DATE"})
    renewal = r.json()["field_key"]
    r = client.post("/contacts/fields", headers=h, json={"label": f"Tier {_uid()}", "field_type": "SELECT",
                                                          "options": ["Gold", "Silver"]})
    tier = r.json()["field_key"]
    emails = [_email(f"bg{i}") for i in range(5)]
    owner = f"{ws['label'].upper()} Owner"
    csv = ("Email,Name,Owner,Phone,Seats,Renewal,Tier\n"
           f"Jane Doe <{emails[0]}>,\"Doe, Jane\",{owner},4155550100.0,\"1,200\",04/03/2026,gold\n"
           f"{emails[1]},Bob Stone,{ws['email']},,(5),,silver\n"
           f"{emails[2]},Cy Young,,,,,\n"
           f"not-an-email,X,,,,,\n"
           f"{emails[3]},Di Lane,,,many,,\n"
           f"{emails[4]},Ed Fox,,,,31/12/2026,Bronze\n")
    up = _upload(client, h, csv.encode())
    import_id = up["import"]["import_id"]
    assert up["import"]["status"] == "UPLOADED" and up["row_count"] == 6
    assert up["suggested_mapping"]["Email"] == "contact.email"
    mapping = {"Email": "contact.email", "Name": "contact.full_name", "Owner": "contact.owner_email",
               "Phone": "contact.phone", "Seats": f"custom.{seats}", "Renewal": f"custom.{renewal}",
               "Tier": f"custom.{tier}"}
    options = {"date_order": "DMY", "default_source": "Expo 2026", "new_list_name": f"Expo {_uid()} {RUN}"}

    r = client.post(f"/imports/{import_id}/dry-run", headers=h, json={"mapping": mapping, "options": options})
    assert r.status_code == 200, r.text
    dry = r.json()
    assert (dry["create"], dry["update"], dry["skip"], dry["fail"]) == (3, 0, 0, 3), dry
    assert sql("SELECT COUNT(*) FROM prospects WHERE email IN :e", e=tuple(emails)) == 0  # nothing written

    r = client.post(f"/imports/{import_id}/start", headers=h, json={"mapping": mapping, "options": options})
    assert r.status_code == 202 and r.json()["status"] == "QUEUED", r.text
    assert client.post(f"/imports/{import_id}/start", headers=h, json={"mapping": mapping}).status_code == 409

    stats = _work(chunk_size=2)
    assert stats["completed"] >= 1
    _check_result(client, ws, import_id, emails, seats, renewal, tier)


def _check_result(client, ws, import_id, emails, seats, renewal, tier):
    h = ws["headers"]
    job = _job(client, h, import_id)
    assert job["status"] == "COMPLETED" and job["progress"] == 100 and job["processed_rows"] == 6
    assert (job["contacts_created"], job["error_count"]) == (3, 3)
    reasons = {e["row"]: e["reason"] for e in job["errors"]}
    assert "Invalid email" in reasons[5] and "is not a number" in reasons[6] and "must be one of" in reasons[7]

    row = sql("SELECT CONCAT_WS('|', first_name, last_name, phone, lead_source, owner_id, "
              "JSON_UNQUOTE(JSON_EXTRACT(custom_fields, :s)), JSON_UNQUOTE(JSON_EXTRACT(custom_fields, :r)), "
              "JSON_UNQUOTE(JSON_EXTRACT(custom_fields, :t))) FROM prospects WHERE email = :e",
              s=f"$.{seats}", r=f"$.{renewal}", t=f"$.{tier}", e=emails[0].lower())
    owner_id = sql("SELECT user_id FROM users WHERE email = :e", e=ws["email"])
    assert row == f"Jane|Doe|4155550100|Expo 2026|{owner_id}|1200|2026-03-04|Gold"
    assert sql("SELECT JSON_EXTRACT(custom_fields, :s) FROM prospects WHERE email = :e",
               s=f"$.{seats}", e=emails[1].lower()) == "-5"
    list_id = sql("SELECT JSON_UNQUOTE(JSON_EXTRACT(options, '$.list_id')) FROM import_jobs WHERE import_id = :i",
                  i=import_id)
    assert sql("SELECT COUNT(*) FROM prospect_list_members WHERE list_id = :l", l=list_id) == 3
    assert sql("SELECT file_path FROM import_jobs WHERE import_id = :i", i=import_id) is not None

    r = client.get(f"/imports/{import_id}/errors.csv", headers=h)
    assert r.status_code == 200 and "not-an-email" in r.content.decode("utf-8-sig")
    hist = client.get("/imports", headers=h).json()
    assert any(i["import_id"] == import_id and i["background"] for i in hist["items"])


def test_stale_running_job_resumes_at_its_checkpoint(client, ws):
    h = ws["headers"]
    emails = [_email(f"rs{i}") for i in range(4)]
    csv = "Email\n" + "".join(f"{e}\n" for e in emails) + f"{emails[0]}\n"
    import_id = _upload(client, h, csv.encode())["import"]["import_id"]
    assert client.post(f"/imports/{import_id}/start", headers=h,
                       json={"mapping": {"Email": "contact.email"}}).status_code == 202
    # A worker took it, committed the first two rows' chunk, and died: RUNNING with an old heartbeat
    sql("UPDATE import_jobs SET status = 'RUNNING', processed_rows = 2, contacts_created = 2, "
        "heartbeat_at = UTC_TIMESTAMP() - INTERVAL 10 MINUTE, claimed_by = 'dead' WHERE import_id = :i", i=import_id)
    _work(chunk_size=2)
    job = _job(client, h, import_id)
    assert job["status"] == "COMPLETED" and job["processed_rows"] == 5
    assert job["contacts_created"] == 4  # 2 before the crash + 2 after; the first two rows were not re-run
    assert sql("SELECT COUNT(*) FROM prospects WHERE email IN :e", e=tuple(e.lower() for e in emails[:2])) == 0
    assert sql("SELECT COUNT(*) FROM prospects WHERE email IN :e", e=tuple(e.lower() for e in emails[2:])) == 2
    assert job["errors"][0]["reason"] == "Duplicate of row 2 in this file"  # remembered across the resume
    assert sql("SELECT attempts FROM import_jobs WHERE import_id = :i", i=import_id) >= 1


def test_a_fresh_running_job_is_not_taken_over(client, ws):
    h = ws["headers"]
    import_id = _upload(client, h, f"Email\n{_email('live')}\n".encode())["import"]["import_id"]
    client.post(f"/imports/{import_id}/start", headers=h, json={"mapping": {"Email": "contact.email"}})
    sql("UPDATE import_jobs SET status = 'RUNNING', heartbeat_at = UTC_TIMESTAMP(), claimed_by = 'alive' "
        "WHERE import_id = :i", i=import_id)
    _work()
    assert _job(client, h, import_id)["status"] == "RUNNING"
    sql("UPDATE import_jobs SET status = 'FAILED' WHERE import_id = :i", i=import_id)
    r = client.post(f"/imports/{import_id}/resume", headers=h)
    assert r.status_code == 202 and r.json()["status"] == "QUEUED"
    _work()
    assert _job(client, h, import_id)["status"] == "COMPLETED"


def test_cancel_before_and_while_running(client, ws, monkeypatch):
    h = ws["headers"]
    import_id = _upload(client, h, f"Email\n{_email('cx')}\n".encode())["import"]["import_id"]
    r = client.post(f"/imports/{import_id}/cancel", headers=h)
    assert r.status_code == 200 and r.json()["status"] == "CANCELLED"
    assert client.post(f"/imports/{import_id}/cancel", headers=h).status_code == 409

    emails = [_email(f"cr{i}") for i in range(6)]
    import_id = _upload(client, h, ("Email\n" + "".join(f"{e}\n" for e in emails)).encode())["import"]["import_id"]
    client.post(f"/imports/{import_id}/start", headers=h, json={"mapping": {"Email": "contact.email"}})
    real = import_service._store_progress

    def cancel_after_first_chunk(job, state, processed):
        real(job, state, processed)
        if job.import_id == import_id:
            job.cancel_requested = True  # what POST /cancel does while the worker runs
    monkeypatch.setattr(import_service, "_store_progress", cancel_after_first_chunk)
    _work(chunk_size=2)
    job = _job(client, h, import_id)
    assert job["status"] == "CANCELLED" and job["processed_rows"] == 2 and job["contacts_created"] == 2
    assert sql("SELECT COUNT(*) FROM prospects WHERE email IN :e", e=tuple(e.lower() for e in emails)) == 2


def _sync(client, h, csv, mapping, options=None):
    return client.post("/imports", headers=h, files={"file": ("c.csv", io.BytesIO(csv.encode()), "text/csv")},
                       data={"mapping": json.dumps(mapping), "options": json.dumps(options or {})})


def test_sync_modes(client, ws):
    h = ws["headers"]
    known, new = _email("known"), _email("new")
    assert _sync(client, h, f"Email,First\n{known},Before\n", {"Email": "contact.email", "First": "contact.first_name"}
                 ).status_code == 201
    csv = f"Email,First\n{known},After\n{new},Newbie\n"
    m = {"Email": "contact.email", "First": "contact.first_name"}
    r = _sync(client, h, csv, m, {"mode": "UPDATE_ONLY"})
    assert r.status_code == 201, r.text
    assert (r.json()["contacts_created"], r.json()["contacts_updated"], r.json()["skipped"]) == (0, 1, 1)
    assert sql("SELECT first_name FROM prospects WHERE email = :e", e=known.lower()) == "After"
    assert sql("SELECT COUNT(*) FROM prospects WHERE email = :e", e=new.lower()) == 0
    r = _sync(client, h, f"Email,First\n{known},Again\n{new},Newbie\n", m, {"mode": "CREATE"})
    assert (r.json()["contacts_created"], r.json()["contacts_updated"], r.json()["skipped"]) == (1, 0, 1)
    assert sql("SELECT first_name FROM prospects WHERE email = :e", e=known.lower()) == "After"
    assert _sync(client, h, csv, m, {"mode": "MERGE"}).status_code == 400


def test_sync_endpoint_queues_a_large_file(client, ws, monkeypatch):
    h = ws["headers"]
    monkeypatch.setattr(import_service, "limits_from_settings",
                        lambda: import_service.ImportLimits(sync_max_rows=2))
    csv = "Email\n" + "".join(f"{_email(f'lg{i}')}\n" for i in range(3))
    r = _sync(client, h, csv, {"Email": "contact.email"})
    assert r.status_code == 202, r.text
    assert r.json()["status"] == "QUEUED" and r.json()["background"] is True
    _work()
    assert _job(client, h, r.json()["import_id"])["contacts_created"] == 3


def test_xlsx_sheet_choice(client, ws):
    h = ws["headers"]
    wb = Workbook()
    wb.active.title = "Notes"
    wb.active.append(["just", "notes"])
    sheet = wb.create_sheet("Leads")
    sheet.append(["Email", "Phone"])
    email = _email("xl")
    sheet.append([email, 4155550199])
    buf = io.BytesIO()
    wb.save(buf)
    up = _upload(client, h, buf.getvalue(), name="book.xlsx", sheet="Leads")
    assert up["sheet_names"] == ["Notes", "Leads"] and up["sheet"] == "Leads" and up["headers"] == ["Email", "Phone"]
    import_id = up["import"]["import_id"]
    client.post(f"/imports/{import_id}/start", headers=h,
                json={"mapping": {"Email": "contact.email", "Phone": "contact.phone"}})
    _work()
    assert sql("SELECT phone FROM prospects WHERE email = :e", e=email.lower()) == "4155550199"


def test_saved_mappings_and_isolation(client, ws, other):
    h, ho = ws["headers"], other["headers"]
    name = f"Trade show {_uid()}"
    r = client.post("/imports/mappings", headers=h, json={"name": name, "mapping": {"Email": "contact.email"},
                                                           "options": {"mode": "CREATE", "date_order": "DMY"}})
    assert r.status_code == 201, r.text
    mapping_id = r.json()["mapping_id"]
    r = client.post("/imports/mappings", headers=h, json={"name": name, "mapping": {"E-mail": "contact.email"}})
    assert r.json()["mapping_id"] == mapping_id and r.json()["mapping"] == {"E-mail": "contact.email"}
    assert any(m["name"] == name for m in client.get("/imports/mappings", headers=h).json())
    assert not any(m["name"] == name for m in client.get("/imports/mappings", headers=ho).json())
    assert client.delete(f"/imports/mappings/{mapping_id}", headers=ho).status_code == 404

    import_id = _upload(client, h, f"Email\n{_email('iso')}\n".encode())["import"]["import_id"]
    assert client.get(f"/imports/{import_id}", headers=ho).status_code == 404
    assert client.post(f"/imports/{import_id}/dry-run", headers=ho, json={"mapping": {}}).status_code == 404
    assert client.post(f"/imports/{import_id}/cancel", headers=ho).status_code == 404
    assert client.get(f"/imports/{import_id}/errors.csv", headers=ho).status_code == 404
    assert client.delete(f"/imports/mappings/{mapping_id}", headers=h).status_code == 200


def test_import_needs_manage_prospects(client, ws):
    from app.core.security import hash_password
    from app.models.user import User
    email = f"noperm-{_uid()}-{RUN}@aimp{RUN}.example.com"
    with SessionLocal() as db:
        db.add(User(tenant_id=ws["tenant_id"], first_name="No", last_name="Perm", email=email,
                    password_hash=hash_password("noperm-pass-1"), role="AGENT", status="ACTIVE", auth_provider="local",
                    email_verified=True, custom_permissions=[]))
        db.commit()
    h = auth(login(client, email, "noperm-pass-1"))
    r = client.post("/imports/uploads", headers=h, files={"file": ("c.csv", io.BytesIO(b"Email\n"), "text/csv")})
    assert r.status_code == 403
    assert client.get("/imports/mappings", headers=h).status_code == 403


def test_retired_address_resolves_to_its_contact(client, ws):
    h = ws["headers"]
    old, new = _email("retired"), _email("current")
    r = client.post("/contacts", headers=h, json={"email": old, "first_name": "Before"})
    assert r.status_code == 201, r.text
    pid = r.json()["prospect_id"]
    r = client.patch(f"/contacts/{pid}", headers=h, json={"email": new})
    assert r.status_code == 200, r.text
    if not sql("SELECT COUNT(*) FROM contact_identifiers WHERE email = :e", e=old.lower()):
        pytest.skip("e-mail history (C-11) is not recorded in this build")
    r = _sync(client, h, f"Email,First\n{old},After\n", {"Email": "contact.email", "First": "contact.first_name"})
    assert r.status_code == 201, r.text
    assert (r.json()["contacts_created"], r.json()["contacts_updated"]) == (0, 1)
    assert sql("SELECT first_name FROM prospects WHERE prospect_id = :p", p=pid) == "After"
    assert sql("SELECT COUNT(*) FROM prospects WHERE email = :e", e=old.lower()) == 0
